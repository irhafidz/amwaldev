"""
Menambahkan slide capaian software Amwal di akhir "Laporan Kemajuan Thesis S2.pptx".

    python build-pptx.py

Gaya mengikuti "2026 MONEV 2 HETI Amwal August 10.pptx": font Inter, navy
0A192F, abu 475569, kartu putih di atas latar motif mosaik, aksen Amwal Emerald
356152. Berkas asli dicadangkan satu kali ke screenshots/_backup/, dan setiap
eksekusi selalu dimulai dari cadangan itu, jadi skrip aman dijalankan ulang
(slide tidak pernah terduplikasi).
"""

from __future__ import annotations

import copy
import re
import shutil
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / "screenshots"
DERIVED = SHOTS / "_pptx"
BACKUP = SHOTS / "_backup" / "Laporan Kemajuan Thesis S2.original.pptx"
TARGET = ROOT / "Laporan Kemajuan Thesis S2.pptx"

EMERALD = RGBColor(0x35, 0x61, 0x52)
EMERALD_DK = RGBColor(0x2B, 0x4F, 0x43)
EMERALD_BG = RGBColor(0xF2, 0xF6, 0xF4)
NAVY = RGBColor(0x0A, 0x19, 0x2F)
SLATE = RGBColor(0x47, 0x55, 0x69)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xFA, 0xCC, 0x15)
AMBER_BG = RGBColor(0xFE, 0xF3, 0xC7)
AMBER_TX = RGBColor(0x92, 0x40, 0x0E)

FONT = "Inter"
FONT_HEAD = "Inter ExtraBold"
FONT_SEMI = "Inter SemiBold"

# Area aman di atas latar: logo di atas (s.d. ~0.65"), motif bawah (mulai ~5.05").
TITLE_Y, SUB_Y = 0.72, 1.15
CONTENT_TOP, CONTENT_BOTTOM = 1.55, 4.95

warnings: list[str] = []

# --------------------------------------------------------------------------
# Perkiraan kecocokan teks (LibreOffice tidak tersedia untuk render)
# --------------------------------------------------------------------------
_FONT_CACHE: dict[tuple[bool, float], ImageFont.FreeTypeFont] = {}
INTER_VS_ARIAL = 1.08  # Inter sedikit lebih lebar dari Arial


def _font(bold: bool, size: float) -> ImageFont.FreeTypeFont:
    key = (bold, size)
    if key not in _FONT_CACHE:
        name = "arialbd.ttf" if bold else "arial.ttf"
        _FONT_CACHE[key] = ImageFont.truetype(f"C:/Windows/Fonts/{name}", int(size * 10))
    return _FONT_CACHE[key]


def est_lines(text: str, size: float, width_in: float, bold: bool = False) -> int:
    font = _font(bold, size)
    max_w = width_in * 72 * 10
    total = 0
    for para in text.split("\n"):
        words = para.split()
        if not words:
            total += 1
            continue
        line, count = "", 1
        for word in words:
            trial = f"{line} {word}".strip()
            if font.getlength(trial) * INTER_VS_ARIAL <= max_w:
                line = trial
            else:
                count += 1
                line = word
        total += count
    return total


def check_fit(label: str, texts: list[tuple[str, float, bool, float]], width_in: float, height_in: float):
    """texts: (teks, ukuran pt, bold, spasi setelah pt)."""
    used = 0.0
    for text, size, bold, after in texts:
        used += est_lines(text, size, width_in, bold) * size * 1.22 / 72 + after / 72
    if used > height_in:
        warnings.append(f"{label}: perkiraan tinggi {used:.2f}in > {height_in:.2f}in")


# --------------------------------------------------------------------------
# Primitif
# --------------------------------------------------------------------------
def set_run(run, size, color=NAVY, font=FONT, bold=False, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, margin=0.0, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    return box, tf


def add_text(slide, x, y, w, h, text, size, color=NAVY, font=FONT, bold=False,
             italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, label=""):
    box, tf = textbox(slide, x, y, w, h, anchor=anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, color, font, bold, italic)
    check_fit(label or text[:30], [(text, size, bold or "Extra" in font, 0)], w, h)
    return box


def card(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.05):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = radius if min(w, h) > 0.6 else 0.2
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def pill(slide, x, y, w, text, fill=EMERALD, color=WHITE, size=8, h=0.24):
    shape = card(slide, x, y, w, h, fill=fill, line=None, radius=0.5)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    set_run(run, size, color, FONT, bold=True)
    return shape


def bullet_paragraph(tf, text, size, first, after=5, color=NAVY, lead=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(after)
    p.line_spacing = 1.05
    if lead:
        r = p.add_run()
        r.text = f"{lead} "
        set_run(r, size, color, FONT, bold=True)
    r = p.add_run()
    r.text = text
    set_run(r, size, color, FONT)
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.17)))
    pPr.set("indent", str(-Inches(0.17)))
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    clr = etree.SubElement(pPr, f"{{{ns}}}buClr")
    etree.SubElement(clr, f"{{{ns}}}srgbClr").set("val", "356152")
    etree.SubElement(pPr, f"{{{ns}}}buFont").set("typeface", "Arial")
    etree.SubElement(pPr, f"{{{ns}}}buChar").set("char", "•")


def bullets(slide, x, y, w, h, items, size=11, after=5, label=""):
    box, tf = textbox(slide, x, y, w, h)
    for i, item in enumerate(items):
        lead, text = item if isinstance(item, tuple) else (None, item)
        bullet_paragraph(tf, text, size, i == 0, after, lead=lead)
    full = [((f"{i[0]} {i[1]}" if isinstance(i, tuple) else i), size, False, after) for i in items]
    check_fit(label or "bullets", full, w - 0.17, h)
    return box


def picture(slide, key, x, y, w, h=None, border=True):
    path = DERIVED / f"{key}.jpg"
    with Image.open(path) as im:
        ratio = im.height / im.width
    h = h or w * ratio
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        pic.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
        pic.line.width = Pt(1)
    return pic, h


def caption_chip(slide, x, y, w, text):
    shape = card(slide, x, y, w, 0.24, fill=NAVY, line=None, radius=0.5)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    set_run(run, 8, WHITE, FONT, bold=True)
    check_fit(f"caption {text}", [(text, 8, True, 0)], w - 0.16, 0.24)


# --------------------------------------------------------------------------
# Slide & kerangka
# --------------------------------------------------------------------------
class Deck:
    def __init__(self, prs: Presentation, bg_part):
        self.prs = prs
        self.layout = next(l for l in prs.slide_layouts if l.name.startswith("Generated"))
        self.bg_part = bg_part

    def new_slide(self):
        slide = self.prs.slides.add_slide(self.layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        rid = slide.part.relate_to(self.bg_part, RT.IMAGE)
        ns_p = "http://schemas.openxmlformats.org/presentationml/2006/main"
        ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        bg = etree.fromstring(
            f'<p:bg xmlns:p="{ns_p}" xmlns:a="{ns_a}" xmlns:r="{ns_r}"><p:bgPr>'
            f'<a:blipFill><a:blip r:embed="{rid}"><a:alphaModFix/></a:blip>'
            f"<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>"
        )
        slide._element.cSld.insert(0, bg)
        return slide

    def header(self, slide, title, subtitle=None):
        add_text(slide, 0.5, TITLE_Y, 9.0, 0.45, title, 24, NAVY, FONT_HEAD, label=f"judul {title}")
        if subtitle:
            add_text(slide, 0.5, SUB_Y, 9.0, 0.3, subtitle, 11, SLATE, FONT, label=f"sub {title}")

    def notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = text


def feature_single(deck, title, subtitle, items, shot, caption, notes=""):
    """Kartu capaian di kiri, satu screenshot desktop di kanan."""
    s = deck.new_slide()
    deck.header(s, title, subtitle)
    card(s, 0.5, 1.6, 4.0, 3.3)
    pill(s, 0.7, 1.78, 0.95, "CAPAIAN")
    bullets(s, 0.7, 2.14, 3.6, 2.65, items, 11, 6, label=f"bullets {title}")
    _, h = picture(s, shot, 4.7, 1.6, 4.8)
    add_text(s, 4.7, 1.6 + h + 0.08, 4.8, 0.3, caption, 9, SLATE, FONT, italic=True, label=f"caption {title}")
    deck.notes(s, notes or " ".join(i if isinstance(i, str) else " ".join(i) for i in items))


def feature_double(deck, title, subtitle, chips, shots, notes=""):
    """Tiga kartu capaian di atas, dua screenshot desktop berdampingan di bawah."""
    s = deck.new_slide()
    deck.header(s, title, subtitle)
    w, gap = 2.9, 0.15
    for i, text in enumerate(chips):
        x = 0.5 + i * (w + gap)
        card(s, x, 1.55, w, 0.8)
        box, tf = textbox(s, x + 0.12, 1.62, w - 0.24, 0.66, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        set_run(r, 10, NAVY, FONT)
        check_fit(f"chip {title} #{i + 1}", [(text, 10, False, 0)], w - 0.24, 0.66)
    iw, slot_h, top = 4.2, 4.2 * 9 / 16, 2.45
    for i, (key, cap) in enumerate(shots):
        x = 0.5 if i == 0 else 9.5 - iw
        with Image.open(DERIVED / f"{key}.jpg") as im:
            h = iw * im.height / im.width
        _, h = picture(s, key, x, top + (slot_h - h) / 2, iw)
        add_text(s, x, top + slot_h + 0.03, iw, 0.2, cap, 9, SLATE, FONT, italic=True, label=f"caption {cap}")
    deck.notes(s, " ".join(chips))


def box(slide, x, y, w, h, title, body, fill=WHITE, tcolor=NAVY, bcolor=SLATE, tsize=10.5, bsize=9):
    card(slide, x, y, w, h, fill=fill)
    b, tf = textbox(slide, x + 0.1, y + 0.07, w - 0.2, h - 0.14, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    set_run(r, tsize, tcolor, FONT, bold=True)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = body
    set_run(r2, bsize, bcolor, FONT)
    check_fit(f"box {title}", [(title, tsize, True, 0), (body, bsize, False, 0)], w - 0.2, h - 0.14)


def arrow(slide, x, y, w=0.2, h=0.25):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = EMERALD
    a.line.fill.background()
    a.shadow.inherit = False


# --------------------------------------------------------------------------
# Gambar turunan (JPEG 1920px; screenshot admin dipotong agar tanpa data pribadi)
# --------------------------------------------------------------------------
def prepare_images():
    DERIVED.mkdir(parents=True, exist_ok=True)
    for src in sorted(SHOTS.glob("[0-9][0-9]-*.png")):
        with Image.open(src) as im:
            im = im.convert("RGB")
            if src.stem == "24-admin-users-rbac":
                # Tabel pengguna memuat email/telepon nyata; hanya bagian atas dipakai.
                im = im.crop((0, 0, im.width, int(im.height * 0.64)))
            scale = 1920 / im.width if im.width > 1920 else 1.0
            if scale < 1:
                im = im.resize((1920, round(im.height * scale)), Image.LANCZOS)
            im.save(DERIVED / f"{src.stem}.jpg", quality=86, optimize=True)


# --------------------------------------------------------------------------
def build():
    if not BACKUP.exists():
        BACKUP.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(TARGET, BACKUP)
        print(f"cadangan: {BACKUP}")
    prs = Presentation(str(BACKUP))
    before = len(prs.slides)

    # Latar mosaik + logo: pakai gambar yang sama dengan slide isi yang sudah ada.
    ref = prs.slides[2]
    rid = re.search(r"<p:bg>.*?r:embed=\"(rId\d+)\"", ref._element.xml, re.S).group(1)
    deck = Deck(prs, ref.part.related_part(rid))
    prepare_images()

    # ---- Ringkasan -------------------------------------------------------
    s = deck.new_slide()
    deck.header(s, "Capaian Software Platform Amwal", "Evaluasi serapan dana & capaian penelitian — Prototype Research TA 2026")
    stats = [("13", "area fitur & integrasi\nyang terdokumentasi"), ("743", "tes otomatis\n(inti 233 + web 510)"),
             ("1.956", "chunk rujukan fiqih\npada indeks RAG"), ("3", "layanan produksi:\nweb, WA bot, RAG")]
    for i, (num, label) in enumerate(stats):
        x = 0.5 + i * 2.3
        card(s, x, 1.65, 2.1, 1.5)
        add_text(s, x + 0.15, 1.75, 1.8, 0.65, num, 32, EMERALD, FONT_HEAD, label=f"stat {num}")
        add_text(s, x + 0.15, 2.42, 1.8, 0.65, label, 10, SLATE, FONT, label=f"statlabel {num}")
    card(s, 0.5, 3.35, 9.0, 1.55, fill=NAVY, line=None)
    box_, tf = textbox(s, 0.75, 3.47, 8.5, 1.35, anchor=MSO_ANCHOR.MIDDLE)
    items = [
        "Satu aplikasi fullstack (TanStack Start + oRPC + Drizzle) dengan kernel domain bersama @amwal/core.",
        "Pembayaran Mayar (QRIS, Virtual Account, E-Wallet), notifikasi WhatsApp, dan asisten AI berbasis RAG sudah terintegrasi.",
        "Audit keamanan C1–C10 dan optimasi performa query telah diselesaikan.",
    ]
    for i, text in enumerate(items):
        bullet_paragraph(tf, text, 11, i == 0, 6, color=RGBColor(0xF8, 0xF9, 0xFA))
    check_fit("ringkasan", [(t, 11, False, 6) for t in items], 8.33, 1.35)
    deck.notes(s, "Ringkasan capaian pengembangan software platform Amwal.")

    # ---- Fitur -----------------------------------------------------------
    feature_double(
        deck, "Auth, RBAC & Profil Pengguna", "Otentikasi, otorisasi berbasis peran, dan manajemen akun",
        ["Login JWT dengan rotasi refresh token; cookie sesi httpOnly.",
         "RBAC 5 peran: Wakif, Nazhir, Admin, Bendahara, Petugas Lapangan.",
         "Reset kata sandi via OTP email dan hardening keamanan (audit C1–C10)."],
        [("23-auth-login", "Halaman masuk"), ("24-admin-users-rbac", "Manajemen akun & hak akses")])

    feature_double(
        deck, "Beranda Publik & Berkah Point", "Beranda donatur, kampanye unggulan, dan gamifikasi",
        ["Beranda menampilkan banner kampanye, ringkasan poin, dan riwayat donasi.",
         "Gamifikasi: absen harian +50 XP, kuis +100 XP, wakaf produktif +200 XP.",
         "Level donatur dan progres XP mendorong kebiasaan berbagi yang istiqamah."],
        [("01-beranda-donatur", "Beranda donatur"), ("25-poin-berkah", "Poin & Keberkahan")])

    # ---- Responsif (tiga tampilan ponsel) --------------------------------
    s = deck.new_slide()
    deck.header(s, "Antarmuka Responsif Mobile-First", "Satu basis kode untuk desktop dan ponsel")
    card(s, 0.5, 1.6, 3.3, 3.3)
    pill(s, 0.7, 1.78, 0.95, "CAPAIAN")
    bullets(s, 0.7, 2.14, 2.95, 2.65, [
        "Navigasi bawah di ponsel, sidebar di desktop.",
        "Widget Tanya Amwal: popup di desktop, halaman penuh di ponsel.",
        "Katalog, form, dan kalkulator disesuaikan untuk layar 390 px.",
    ], 11, 6, label="bullets responsif")
    ph = 2.85
    for i, (key, cap) in enumerate([("21-beranda-mobile", "Beranda"), ("05-chatbot-mobile", "Tanya Amwal"), ("22-wakaf-mobile", "Katalog wakaf")]):
        x = 4.15 + i * 1.8
        pic, _ = picture(s, key, x, 1.6, ph * 390 / 844, ph)
        add_text(s, x - 0.2, 1.6 + ph + 0.08, ph * 390 / 844 + 0.4, 0.25, cap, 9, SLATE, FONT, italic=True,
                 align=PP_ALIGN.CENTER, label=f"cap {cap}")
    deck.notes(s, "Tampilan mobile 390 px.")

    feature_double(
        deck, "Layanan Wakaf Digital", "Katalog program dan transparansi laporan Nazhir",
        ["Katalog program berkategori; detail progres, target, dan pilihan nominal wakaf.",
         "Ledger pokok wakaf, hasil kelola, dan penarikan dana dengan persetujuan.",
         "Nazhir mengunggah laporan progres dan bukti; sertifikat wakaf dapat diverifikasi."],
        [("06-wakaf-katalog", "Katalog program wakaf"), ("07-wakaf-detail", "Detail program & donasi")])

    feature_double(
        deck, "Layanan Zakat & Kalkulator", "Kalkulator otomatis syariah: fitrah dan maal",
        ["Zakat fitrah dari tarif terkonfigurasi; maal: profesi, perusahaan, perdagangan, emas.",
         "Nisab 85 gram emas dihitung dari harga emas terkini (riwayat harga tersimpan).",
         "Penyaluran ke 8 asnaf dan pencatatan mustahik oleh admin dan amil."],
        [("08-zakat-kalkulator-fitrah", "Kalkulator zakat fitrah"), ("09-zakat-kalkulator-maal", "Kalkulator zakat maal")])

    feature_single(
        deck, "Layanan Infaq & Sedekah", "Multi-akad dengan alokasi program yang jelas",
        [("5 jenis:", "Infaq Subuh, Infaq Jumat, Infaq Umum, Sedekah Yatim, Sedekah Operasional."),
         "Nominal cepat atau bebas; ringkasan biaya transparan sebelum bayar.",
         "Pembayaran via Mayar: QRIS, Virtual Account, atau E-Wallet.",
         "Notifikasi terima kasih otomatis lewat WhatsApp."],
        "10-infaq-sedekah", "Gambar: halaman Infaq & Sedekah Berkah")

    feature_single(
        deck, "Layanan Qurban Digital Terpadu", "Pemesanan hingga laporan distribusi",
        ["Pemesanan kambing dan sapi, termasuk patungan kolektif (1/7).",
         "Batch hewan dengan slot dan status ketersediaan (Penuh / Tersisa).",
         "Pilihan pasrah atau ambil sendiri; laporan distribusi daging.",
         "Setoran petugas lapangan dan pengelolaan RPH di sisi admin."],
        "11-qurban", "Gambar: katalog paket qurban Idul Adha")

    feature_double(
        deck, "Edukasi Syariah & Kuis Interaktif", "Berita, artikel, dan kuis yang terhubung ke poin",
        ["Artikel dan video edukasi berkategori, lengkap dengan suka dan komentar.",
         "Kuis pemahaman per artikel memberi +100 XP Berkah.",
         "CMS admin untuk berita dan edukasi; konten dirender aman (anti-XSS)."],
        [("13-edukasi", "Katalog edukasi"), ("14-edukasi-kuis", "Artikel & kuis pemahaman")])

    feature_double(
        deck, "Fitur AI Assistant ZISWAF", "RAG pipeline: Jina + FAISS (1.956 chunk) + DeepSeek",
        ["Pertanyaan dicari di indeks FAISS 1.956 chunk, dire-rank Jina, dijawab DeepSeek.",
         "Jawaban memuat sumber rujukan (dokumen, bab, halaman); menolak bila tidak ditemukan.",
         "Widget melayang di dasbor dan halaman /tanya-ai; timeout 70 dtk, retry 429/503."],
        [("03-chatbot-widget-popup", "Widget Tanya Amwal (popup)"), ("04-chatbot-halaman-penuh", "Halaman penuh /tanya-ai")])

    feature_double(
        deck, "Admin Dashboard & Kampanye", "Ringkasan operasional dan manajemen program",
        ["Ringkasan dari transaksi lunas: penghimpunan, donatur aktif, dan program berjalan.",
         "Analitik RFM-D untuk segmentasi donatur serta heatmap retensi.",
         "Penerbitan program wakaf, progres fisik, dan verifikasi Nazhir."],
        [("15-admin-overview", "Ringkasan operasional"), ("16-admin-wakaf", "Manajemen program wakaf")])

    feature_double(
        deck, "Admin Transaksi & Mutasi Keuangan", "Transparansi transaksi dan modul ZISWAF",
        ["Tabel transparansi transaksi terverifikasi dan log realisasi dana.",
         "Laporan Transparansi dapat dicetak PDF; snapshot alokasi per program.",
         "Modul ZISWAF: infaq, zakat 8 asnaf, dan qurban dikelola terpisah."],
        [("19-admin-transaksi-laporan", "Transparansi & buku besar"), ("17-admin-ziswaf-zakat", "Modul zakat & 8 asnaf")])

    # ---- Payment gateway (diagram alur) ----------------------------------
    s = deck.new_slide()
    deck.header(s, "Integrasi Payment Gateway Mayar", "QRIS, Virtual Account, dan E-Wallet dengan webhook dan cek status")
    steps = [("1  Pesanan", "ID urut atomik\nper modul"), ("2  Mayar", "QRIS, VA bank,\natau E-Wallet"),
             ("3  Konfirmasi", "Webhook atau\ncek status langsung"), ("4  Settlement", "Idempoten:\nstatus LUNAS"),
             ("5  Ledger", "Kuitansi dan\nnotifikasi WA")]
    bw, bg_ = 1.6, 0.25
    for i, (t, b) in enumerate(steps):
        x = 0.5 + i * (bw + bg_)
        box(s, x, 1.6, bw, 1.0, t, b, fill=EMERALD_BG if i % 2 == 0 else WHITE)
        if i < len(steps) - 1:
            arrow(s, x + bw + 0.025, 2.0)
    card(s, 0.5, 2.85, 4.4, 2.05)
    pill(s, 0.7, 3.0, 0.95, "CAPAIAN")
    bullets(s, 0.7, 3.34, 4.05, 1.5, [
        "Donatur memilih QRIS, Virtual Account (7 bank), atau E-Wallet (4 aplikasi).",
        "Halaman status menanyakan Mayar langsung, jadi lunas tanpa menunggu webhook.",
    ], 10.5, 5, label="mayar capaian")
    card(s, 5.1, 2.85, 4.4, 2.05)
    pill(s, 5.3, 3.0, 0.95, "PENGAMAN", fill=NAVY)
    bullets(s, 5.3, 3.34, 4.05, 1.5, [
        "Token webhook dibandingkan constant-time; status diambil ulang dari API.",
        "Pembaruan hanya bila belum LUNAS, sehingga tidak ada hitung ganda.",
        "Job rekonsiliasi memproses transaksi PENDING yang konfirmasinya tidak sampai.",
    ], 10.5, 5, label="mayar pengaman")
    deck.notes(s, "Alur pembayaran Mayar: pesanan, pembayaran, konfirmasi, settlement idempoten, ledger.")

    feature_single(
        deck, "Integrasi Notifikasi WhatsApp", "Baileys engine untuk notifikasi transaksi dan akta ikrar wakaf",
        ["Mesin Baileys berjalan sebagai layanan terpisah (wa-bot), tanpa rute publik.",
         "Pairing perangkat lewat QR di panel admin dengan status koneksi langsung.",
         "Templat terima kasih wakaf, zakat, dan infaq; nomor HP dinormalisasi.",
         "Panggilan antarlayanan memakai rahasia internal pada header khusus."],
        "18-admin-whatsapp-gateway", "Gambar: panel admin, pairing QR Baileys (sesi uji)")

    # ---- Arsitektur deployment -------------------------------------------
    s = deck.new_slide()
    deck.header(s, "Arsitektur Deployment", "Dokploy, VPS Contabo, Docker Swarm, dan MinIO S3")
    card(s, 0.5, 1.55, 9.0, 0.4, fill=NAVY, line=None)
    add_text(s, 0.5, 1.55, 9.0, 0.4, "Pengguna  →  amwal.id (HTTPS)  →  Traefik reverse proxy (Dokploy)", 10.5,
             WHITE, FONT, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, label="arsitektur atas")
    svc = [("Web App", "TanStack Start + oRPC + Drizzle"), ("WA Bot", "Baileys, hanya jaringan internal"),
           ("RAG Service", "FastAPI, FAISS, Jina, DeepSeek"), ("PostgreSQL", "Skema tunggal @amwal/core, 19 migrasi"),
           ("MinIO (S3)", "Kuitansi dan dokumen")]
    sw, sg = 1.7, 0.125
    for i, (t, b) in enumerate(svc):
        box(s, 0.5 + i * (sw + sg), 2.15, sw, 1.3, t, b)
    card(s, 0.5, 3.65, 9.0, 0.45, fill=EMERALD, line=None)
    add_text(s, 0.5, 3.65, 9.0, 0.45, "Docker Compose + Docker Swarm di jaringan overlay dokploy-network, dikelola Dokploy", 10.5,
             WHITE, FONT, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, label="arsitektur tengah")
    card(s, 0.5, 4.25, 9.0, 0.65, fill=EMERALD_BG)
    add_text(s, 0.65, 4.25, 8.7, 0.65, "VPS Contabo  ·  deploy dari branch main  ·  quality gate: Biome, type-check, dan 743 tes otomatis",
             10.5, NAVY, FONT, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, label="arsitektur bawah")
    deck.notes(s, "Arsitektur deployment produksi platform Amwal.")

    # ---- Realisasi anggaran ----------------------------------------------
    s = deck.new_slide()
    deck.header(s, "Realisasi Anggaran", "Pembelanjaan infrastruktur dan layanan AI untuk Prototype Research")
    rows = [
        ("1", "Domain amwal.id", "Rumahweb", "Alamat produksi platform"),
        ("2", "VPS Contabo", "Contabo", "Server produksi: web, WA bot, RAG, database, MinIO"),
        ("3", "Token AI DeepSeek", "DeepSeek", "Model bahasa penyusun jawaban asisten AI"),
        ("4", "Token AI Jina", "Jina AI", "Embedding dan reranker untuk pencarian rujukan"),
    ]
    heads = ["No", "Item", "Penyedia", "Fungsi", "Nominal (Rp)"]
    widths = [0.45, 1.9, 1.3, 3.65, 1.7]
    nrows = len(rows) + 2
    shape = s.shapes.add_table(nrows, 5, Inches(0.5), Inches(1.6), Inches(sum(widths)), Inches(3.1))
    tbl = shape.table
    # Tanpa gaya tabel bawaan: semua sel diberi warna eksplisit.
    for j, wcol in enumerate(widths):
        tbl.columns[j].width = Inches(wcol)
    tbl.rows[0].height = Inches(0.4)
    for i in range(1, nrows - 1):
        tbl.rows[i].height = Inches(0.6)
    tbl.rows[nrows - 1].height = Inches(0.4)

    def cell(i, j, text, fill, color=NAVY, bold=False, size=10, align=PP_ALIGN.LEFT):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = fill
        c.margin_left = c.margin_right = Inches(0.08)
        c.margin_top = c.margin_bottom = Inches(0.04)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        set_run(r, size, color, FONT, bold=bold)
        c.text_frame.word_wrap = True
        span = sum(widths[:4]) if (i == nrows - 1 and j == 0) else widths[j]
        check_fit(f"sel[{i},{j}] {text[:20]}", [(text, size, bold, 0)], span - 0.16, tbl.rows[i].height / 914400 - 0.08)

    for j, h in enumerate(heads):
        cell(0, j, h, NAVY, WHITE, True, 10, PP_ALIGN.CENTER if j in (0, 4) else PP_ALIGN.LEFT)
    for i, row in enumerate(rows, start=1):
        band = WHITE if i % 2 else EMERALD_BG
        for j, text in enumerate(row):
            cell(i, j, text, band, bold=(j == 1), align=PP_ALIGN.CENTER if j == 0 else PP_ALIGN.LEFT)
        cell(i, 4, "Isi sesuai bukti bayar", AMBER_BG, AMBER_TX, True, 9, PP_ALIGN.CENTER)
    last = nrows - 1
    tbl.cell(last, 0).merge(tbl.cell(last, 3))
    cell(last, 0, "Total realisasi", EMERALD, WHITE, True, 10.5, PP_ALIGN.RIGHT)
    cell(last, 4, "Isi sesuai bukti bayar", AMBER_BG, AMBER_TX, True, 9, PP_ALIGN.CENTER)
    deck.notes(s, "Nominal pada kolom kanan harus diisi dari bukti pembayaran sebelum laporan dikirim.")

    # ---- Roadmap ---------------------------------------------------------
    s = deck.new_slide()
    deck.header(s, "Roadmap & Rencana Tindak Lanjut", "Dari prototype menuju uji coba lapangan dan integrasi nasional")
    phases = [
        ("1", "Uji coba mitra lapangan", [
            "Onboarding Nazhir dan petugas lapangan mitra.",
            "Uji end-to-end wakaf, zakat, infaq, dan qurban dengan transaksi nyata.",
            "Kumpulkan umpan balik UX sebelum rilis publik."]),
        ("2", "Evaluasi RAG", [
            "Susun set pertanyaan uji fiqih bersama asatidz.",
            "Ukur akurasi jawaban, kelengkapan rujukan, dan tingkat penolakan.",
            "Perluas korpus dan sesuaikan retrieval; selaras skripsi chatbot."]),
        ("3", "Integrasi API BWI", [
            "Petakan data program wakaf ke format pelaporan BWI.",
            "Sinkronkan pelaporan Nazhir dan sertifikat wakaf.",
            "Uji koneksi API dan kebijakan keamanan data."]),
    ]
    pw, pg = 2.9, 0.15
    for i, (num, title, pts) in enumerate(phases):
        x = 0.5 + i * (pw + pg)
        card(s, x, 1.6, pw, 2.3)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.15), Inches(1.72), Inches(0.36), Inches(0.36))
        circ.fill.solid()
        circ.fill.fore_color.rgb = EMERALD
        circ.line.fill.background()
        circ.shadow.inherit = False
        ctf = circ.text_frame
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = num
        set_run(cr, 11, WHITE, FONT, bold=True)
        add_text(s, x + 0.6, 1.72, pw - 0.7, 0.36, title, 11.5, NAVY, FONT_SEMI, bold=True,
                 anchor=MSO_ANCHOR.MIDDLE, label=f"fase {title}")
        bullets(s, x + 0.15, 2.22, pw - 0.3, 1.6, pts, 10, 5, label=f"fase {num}")
    card(s, 0.5, 4.1, 9.0, 0.7, fill=NAVY, line=None)
    add_text(s, 0.7, 4.1, 8.6, 0.7, "Penguatan produksi: kunci API Mayar produksi  ·  pairing WhatsApp produksi  ·  pemantauan kuota token AI dan domain",
             10, RGBColor(0xF8, 0xF9, 0xFA), FONT, bold=True, anchor=MSO_ANCHOR.MIDDLE, label="roadmap strip")
    deck.notes(s, "Rencana tindak lanjut setelah Monev.")

    prs.save(str(TARGET))
    print(f"slide: {before} -> {len(prs.slides)} (ditambah {len(prs.slides) - before})")
    if warnings:
        print("\nPERINGATAN kecocokan teks:")
        for w in warnings:
            print(" -", w)
        return 1
    print("kecocokan teks: OK")
    return 0


if __name__ == "__main__":
    sys.exit(build())
