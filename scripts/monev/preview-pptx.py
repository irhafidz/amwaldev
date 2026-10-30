"""
Pratinjau kasar slide PPTX tanpa LibreOffice: menggambar latar, bentuk, gambar,
tabel, dan teks (Arial diperlebar 8% sebagai pengganti Inter) ke PNG.
Dipakai hanya untuk memeriksa tumpang tindih/luapan; bukan render PowerPoint.

    python preview-pptx.py [slide_awal] [slide_akhir]
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / "Laporan Kemajuan Thesis S2.pptx"
OUT = ROOT / "screenshots" / "_preview"
SCALE = 160  # px per inci
FACTOR = 1.08


def font(size, bold):
    return ImageFont.truetype(f"C:/Windows/Fonts/{'arialbd' if bold else 'arial'}.ttf", max(6, round(size * SCALE / 72)))


def wrap(draw, text, fnt, width):
    lines = []
    for para in text.split("\n"):
        line = ""
        for word in para.split():
            trial = f"{line} {word}".strip()
            if draw.textlength(trial, font=fnt) * FACTOR <= width:
                line = trial
            else:
                if line:
                    lines.append(line)
                line = word
        lines.append(line)
    return lines


def rgb(color_fmt, default=None):
    try:
        return tuple(color_fmt.rgb)
    except Exception:
        return default


def draw_text_frame(draw, tf, x, y, w, h):
    ml = (tf.margin_left or 0) / 914400
    mt = (tf.margin_top or 0) / 914400
    cx, cy, cw = x + ml * SCALE, y + mt * SCALE, w - 2 * ml * SCALE
    rows = []
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        size = runs[0].font.size.pt if runs[0].font.size else 11
        bold = any(r.font.bold or "Extra" in (r.font.name or "") or "Semi" in (r.font.name or "") for r in runs)
        color = rgb(runs[0].font.color, (10, 25, 47))
        text = "".join(r.text for r in runs)
        indent = 0.17 * SCALE if p._p.pPr is not None and p._p.pPr.get("marL") else 0
        fnt = font(size, bold)
        lines = wrap(draw, text, fnt, cw - indent)
        after = (p.space_after.pt if p.space_after else 0) * SCALE / 72
        rows.append((lines, fnt, color, size, indent, after, p.alignment, bool(indent)))
    total = sum(len(l) * s * 1.22 * SCALE / 72 + a for l, _, _, s, _, a, _, _ in rows)
    anchor = tf.vertical_anchor
    top = cy
    if anchor is not None and int(anchor) == 3:  # MIDDLE
        top = y + (h - total) / 2
    pos = top
    for lines, fnt, color, size, indent, after, align, bullet in rows:
        for k, line in enumerate(lines):
            tw = draw.textlength(line, font=fnt) * FACTOR
            lx = cx + indent
            if align is not None and int(align) == 2:  # CENTER
                lx = cx + (cw - tw) / 2
            elif align is not None and int(align) == 3:  # RIGHT
                lx = cx + cw - tw
            draw.text((lx, pos), line, font=fnt, fill=color)
            if bullet and k == 0:
                draw.text((cx, pos), "•", font=fnt, fill=(53, 97, 82))
            pos += size * 1.22 * SCALE / 72
        pos += after
    return pos > y + h + 2  # luapan


def render(slide, idx, bg):
    W, H = round(10 * SCALE), round(5.625 * SCALE)
    img = bg.resize((W, H)).convert("RGB")
    d = ImageDraw.Draw(img)
    over = []

    def walk(shapes):
        for sh in shapes:
            x, y = Emu(sh.left).inches * SCALE, Emu(sh.top).inches * SCALE
            w, h = Emu(sh.width).inches * SCALE, Emu(sh.height).inches * SCALE
            if sh.shape_type == 13:
                im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB").resize((round(w), round(h)))
                img.paste(im, (round(x), round(y)))
                d.rectangle([x, y, x + w, y + h], outline=(203, 213, 225))
                continue
            if sh.has_table if hasattr(sh, "has_table") else False:
                yy = y
                for r in sh.table.rows:
                    rh = r.height / 914400 * SCALE
                    xx = x
                    for ci, c in enumerate(r.cells):
                        cw = sh.table.columns[ci].width / 914400 * SCALE
                        if c.is_spanned:
                            xx += cw
                            continue
                        span = c.span_width if c.is_merge_origin else 1
                        cw_total = sum(sh.table.columns[ci + k].width for k in range(span)) / 914400 * SCALE
                        fill = rgb(c.fill.fore_color, (255, 255, 255))
                        d.rectangle([xx, yy, xx + cw_total, yy + rh], fill=fill, outline=(226, 232, 240))
                        if draw_text_frame(d, c.text_frame, xx, yy, cw_total, rh):
                            over.append(f"tabel sel {ci}")
                        xx += cw
                    yy += rh
                continue
            if sh.shape_type in (1, 5) or sh.shape_type is None:
                try:
                    fill = rgb(sh.fill.fore_color) if sh.fill.type == 1 else None
                except Exception:
                    fill = None
                try:
                    line = rgb(sh.line.color) if sh.line.fill.type == 1 else None
                except Exception:
                    line = None
                rr = 0.5 if sh.auto_shape_type in (5, 9) and min(w, h) < 60 else 0
                if fill or line:
                    if sh.auto_shape_type == 9:
                        d.ellipse([x, y, x + w, y + h], fill=fill, outline=line)
                    elif sh.auto_shape_type == 33:  # RIGHT_ARROW
                        d.polygon([(x, y + h * .25), (x + w * .6, y + h * .25), (x + w * .6, y), (x + w, y + h / 2),
                                   (x + w * .6, y + h), (x + w * .6, y + h * .75), (x, y + h * .75)], fill=fill)
                    else:
                        d.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) * (0.5 if rr else 0.05), fill=fill, outline=line)
            if sh.has_text_frame and sh.text_frame.text.strip():
                if draw_text_frame(d, sh.text_frame, x, y, w, h):
                    over.append(sh.text_frame.text[:40].replace("\n", " "))
    walk(slide.shapes)
    # garis aman: logo atas (0.65") dan motif bawah (5.05")
    d.line([(0, 0.65 * SCALE), (W, 0.65 * SCALE)], fill=(255, 0, 0), width=1)
    d.line([(0, 5.05 * SCALE), (W, 5.05 * SCALE)], fill=(255, 0, 0), width=1)
    OUT.mkdir(parents=True, exist_ok=True)
    img.save(OUT / f"slide-{idx:02d}.png")
    return over


if __name__ == "__main__":
    prs = Presentation(str(TARGET))
    first = int(sys.argv[1]) if len(sys.argv) > 1 else 6
    last = int(sys.argv[2]) if len(sys.argv) > 2 else len(prs.slides)
    import re
    ref = prs.slides[2]
    rid = re.search(r"<p:bg>.*?r:embed=\"(rId\d+)\"", ref._element.xml, re.S).group(1)
    bg = Image.open(io.BytesIO(ref.part.related_part(rid).blob))
    for i in range(first, last + 1):
        problems = render(prs.slides[i - 1], i, bg)
        print(f"slide {i}: {'LUAPAN -> ' + '; '.join(problems) if problems else 'ok'}")
